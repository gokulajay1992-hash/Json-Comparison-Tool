import streamlit as st
import oracledb
import json
import io
from datetime import datetime, date, time
import openpyxl
from openpyxl.styles import PatternFill, Font

st.set_page_config(page_title="JSON Comparison Tool", layout="wide")

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

def load_config():
    with open("config.json") as f:
        return json.load(f)

CFG = load_config()
COL = CFG["columns"]
EXCLUDE_FIELDS = set(CFG.get("exclude_fields", []))

# ---------------------------------------------------------------------------
# Oracle
# ---------------------------------------------------------------------------

def get_connection(host, port, service, username, password):
    dsn = oracledb.makedsn(host, int(port), service_name=service)
    try:
        conn = oracledb.connect(user=username, password=password, dsn=dsn)
        return conn
    except oracledb.DatabaseError as e:
        error, = e.args
        raise ConnectionError(
            f"Oracle connection failed.\n"
            f"  Host:    {host}:{port}\n"
            f"  Service: {service}\n"
            f"  User:    {username}\n"
            f"  Detail:  {error.message}"
        )

# ---------------------------------------------------------------------------
# JSON diff engine
# ---------------------------------------------------------------------------

def _fmt(val):
    if val is None:
        return ""
    if isinstance(val, (dict, list)):
        return json.dumps(val)
    return str(val)

def _is_empty(obj):
    if isinstance(obj, (dict, list)):
        return len(obj) == 0
    return False

def strip_excluded(obj, exclude=EXCLUDE_FIELDS):
    """Recursively remove excluded fields and any empty dicts/lists they leave behind."""
    if isinstance(obj, dict):
        result = {}
        for k, v in obj.items():
            if k in exclude:
                continue
            cleaned = strip_excluded(v, exclude)
            if not _is_empty(cleaned):
                result[k] = cleaned
        return result
    elif isinstance(obj, list):
        cleaned = [strip_excluded(item, exclude) for item in obj]
        return [item for item in cleaned if not _is_empty(item)]
    return obj

def deep_diff(a, b, path=""):
    diffs = []
    if type(a) != type(b):
        diffs.append({"path": path or "root", "old": _fmt(a), "new": _fmt(b), "type": "changed"})
        return diffs
    if isinstance(a, dict):
        for key in set(a) | set(b):
            p = f"{path}.{key}" if path else key
            if key not in a:
                diffs.append({"path": p, "old": "", "new": _fmt(b[key]), "type": "added"})
            elif key not in b:
                diffs.append({"path": p, "old": _fmt(a[key]), "new": "", "type": "removed"})
            else:
                diffs.extend(deep_diff(a[key], b[key], p))
    elif isinstance(a, list):
        for i in range(max(len(a), len(b))):
            p = f"{path}[{i}]"
            if i >= len(a):
                diffs.append({"path": p, "old": "", "new": _fmt(b[i]), "type": "added"})
            elif i >= len(b):
                diffs.append({"path": p, "old": _fmt(a[i]), "new": "", "type": "removed"})
            else:
                diffs.extend(deep_diff(a[i], b[i], p))
    else:
        if a != b:
            diffs.append({"path": path or "root", "old": _fmt(a), "new": _fmt(b), "type": "changed"})
    return diffs

def parse_json(val):
    if val is None:
        return None
    if isinstance(val, str):
        try:
            return json.loads(val)
        except Exception:
            return val
    return val

# ---------------------------------------------------------------------------
# Query
# ---------------------------------------------------------------------------

def run_query(host, port, service, username, password, date_from, date_to, primary_id_val=None, id1_val=None, id2_val=None):
    extra_ids   = COL.get("extra_ids", [])
    json_sets   = COL["json_sets"]
    primary_id  = COL["primary_id"]
    dt_col      = COL["datetime"]

    id1_col = extra_ids[0] if len(extra_ids) > 0 else None
    id2_col = extra_ids[1] if len(extra_ids) > 1 else None

    select_cols = [primary_id, dt_col] + extra_ids
    for s in json_sets:
        for side in (s["left"], s["right"]):
            if side not in select_cols:
                select_cols.append(side)

    col_list  = ", ".join(select_cols)
    where     = f"{dt_col} >= :date_from AND {dt_col} <= :date_to"
    bind_vars = {
        "date_from": datetime.strptime(date_from, "%Y-%m-%d %H:%M:%S"),
        "date_to":   datetime.strptime(date_to,   "%Y-%m-%d %H:%M:%S"),
    }

    if primary_id_val:
        where += f" AND {primary_id} = :primary_id_val"
        bind_vars["primary_id_val"] = primary_id_val

    if id1_val and id1_col:
        where += f" AND {id1_col} = :id1_val"
        bind_vars["id1_val"] = id1_val

    if id2_val and id2_col:
        where += f" AND {id2_col} = :id2_val"
        bind_vars["id2_val"] = id2_val

    sql  = f"SELECT {col_list} FROM {CFG['table_name']} WHERE {where} ORDER BY {dt_col}"

    conn = get_connection(host, port, service, username, password)

    def clob_to_str(cursor, name, default_type, size, precision, scale):
        if default_type in (oracledb.DB_TYPE_CLOB, oracledb.DB_TYPE_NCLOB):
            return cursor.var(oracledb.DB_TYPE_LONG, arraysize=cursor.arraysize)

    conn.outputtypehandler = clob_to_str
    cursor = conn.cursor()
    cursor.execute(sql, bind_vars)
    rows   = cursor.fetchall()
    cursor.close()
    conn.close()

    results = []
    for row in rows:
        row_map = {c: row[i] for i, c in enumerate(select_cols)}
        record  = {
            "id":       str(row_map.get(primary_id, "")),
            "dat_time": str(row_map.get(dt_col, "")),
            "extra":    {k: str(row_map.get(k, "") or "") for k in extra_ids},
            "sets":     [],
        }
        for s in json_sets:
            left  = parse_json(row_map.get(s["left"]))
            right = parse_json(row_map.get(s["right"]))
            if left is not None and right is not None:
                # Auto-detect which side has the wrapper and unwrap it
                unwrap_keys = [k for k in [s.get("left_unwrap",""), s.get("right_unwrap","")] if k]
                for uw in unwrap_keys:
                    if isinstance(left, dict) and uw in left:
                        left = left[uw]
                    if isinstance(right, dict) and uw in right:
                        right = right[uw]
                diffs = deep_diff(strip_excluded(left), strip_excluded(right))
            else:
                diffs = []
            record["sets"].append({
                "label":     s["label"],
                "left_col":  s["left"],
                "right_col": s["right"],
                "left_label":  s.get("left_label", s["left"]),
                "right_label": s.get("right_label", s["right"]),
                "diffs":     diffs,
                "has_diff":  len(diffs) > 0,
            })
        results.append(record)

    return results

# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def build_excel(results):
    extra_ids = COL.get("extra_ids", [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "JSON Diff"

    hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hdr_font = Font(color="FFFFFF", bold=True)
    fills    = {
        "added":   PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "removed": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "changed": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    }

    base_headers = [COL["primary_id"]] + extra_ids + [COL["datetime"]]
    headers = base_headers + ["Comparison Set", "Left Column", "Right Column",
                               "Path", "Difference Type", "Old Value", "New Value"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font

    for rec in results:
        base = [rec["id"]] + [rec["extra"].get(k, "") for k in extra_ids] + [rec["dat_time"]]
        for s in rec["sets"]:
            if not s["diffs"]:
                ws.append(base + [s["label"], s["left_label"], s["right_label"],
                                   "(no differences)", "", "", ""])
            for d in s["diffs"]:
                ws.append(base + [s["label"], s["left_label"], s["right_label"],
                                   d["path"], d["type"], d["old"], d["new"]])
                if d["type"] in fills:
                    for cell in ws[ws.max_row]:
                        cell.fill = fills[d["type"]]

    for col_cells in ws.columns:
        width = max(len(str(c.value or "")) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(width + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# ---------------------------------------------------------------------------
# UI helpers
# ---------------------------------------------------------------------------

DIFF_COLORS = {"added": "#c6efce", "removed": "#ffc7ce", "changed": "#ffeb9c"}
DIFF_TEXT   = {"added": "#276221", "removed": "#8b0000",  "changed": "#7d5800"}

def diff_badge(dtype):
    bg = DIFF_COLORS.get(dtype, "#eee")
    fg = DIFF_TEXT.get(dtype, "#000")
    return f'<span style="background:{bg};color:{fg};padding:1px 7px;border-radius:3px;font-size:.75rem;font-weight:700">{dtype.upper()}</span>'

def render_diffs(diffs):
    for d in diffs:
        bg = DIFF_COLORS.get(d["type"], "#fafafa")
        border = {"added": "#4caf50", "removed": "#f44336", "changed": "#ff9800"}.get(d["type"], "#ccc")
        old_html = f'<span style="background:#ffd7d7;padding:1px 6px;border-radius:3px;font-family:monospace;font-size:.8rem">{d["old"]}</span>' if d["old"] else ""
        new_html = f'<span style="background:#d7ffd7;padding:1px 6px;border-radius:3px;font-family:monospace;font-size:.8rem">{d["new"]}</span>' if d["new"] else ""
        arrow    = " &rarr; " if old_html and new_html else ""
        st.markdown(
            f'<div style="background:{bg};border-left:3px solid {border};padding:6px 10px;'
            f'border-radius:4px;margin-bottom:4px">'
            f'{diff_badge(d["type"])} &nbsp;<code style="font-size:.8rem">{d["path"]}</code><br/>'
            f'<div style="margin-top:4px">{old_html}{arrow}{new_html}</div>'
            f'</div>',
            unsafe_allow_html=True,
        )

# ---------------------------------------------------------------------------
# Main app
# ---------------------------------------------------------------------------

st.title("JSON Comparison Tool")

# ── Sidebar ────────────────────────────────────────────────────────────────
with st.sidebar:
    st.header("DB Connection")
    host     = st.text_input("Host", placeholder="db.host.com")
    port     = st.text_input("Port", value="1521")
    service  = st.text_input("Service Name / SID", placeholder="ORCL")
    username = st.text_input("Username")
    password = st.text_input("Password", type="password")

    st.divider()
    st.header("Date & Time Range")
    col1, col2 = st.columns(2)
    with col1:
        from_date = st.date_input("From date", value=date.today())
        from_time = st.time_input("From time", value=time(0, 0, 0))
    with col2:
        to_date = st.date_input("To date", value=date.today())
        to_time = st.time_input("To time", value=time(23, 59, 59))

    st.divider()
    _extra = COL.get("extra_ids", [])
    st.header("Filter by ID (optional)")
    primary_id_label = COL.get("primary_id", "ID")
    id1_label = _extra[0] if len(_extra) > 0 else "ID1"
    id2_label = _extra[1] if len(_extra) > 1 else "ID2"
    primary_id_val = st.text_input(primary_id_label, placeholder="Leave blank to fetch all")
    id1_val = st.text_input(id1_label, placeholder="Leave blank to fetch all")
    id2_val = st.text_input(id2_label, placeholder="Leave blank to fetch all")

    test_btn = st.button("Test Connection", use_container_width=True)
    run_btn  = st.button("Run Comparison", type="primary", use_container_width=True)

# ── Test connection ────────────────────────────────────────────────────────
if test_btn:
    if not all([host, service, username, password]):
        st.error("Please fill in all DB connection fields.")
    else:
        with st.spinner("Connecting..."):
            try:
                conn = get_connection(host, port, service, username, password)
                conn.close()
                st.success(f"Connected successfully to **{service}** on **{host}:{port}**")
            except ConnectionError as e:
                st.error(str(e))
            except Exception as e:
                st.error(f"Unexpected error: {e}")

# ── Run query ──────────────────────────────────────────────────────────────
if run_btn:
    if not all([host, service, username, password]):
        st.error("Please fill in all DB connection fields.")
    else:
        date_from = datetime.combine(from_date, from_time).strftime("%Y-%m-%d %H:%M:%S")
        date_to   = datetime.combine(to_date,   to_time).strftime("%Y-%m-%d %H:%M:%S")
        with st.spinner("Querying Oracle DB..."):
            try:
                st.session_state["results"] = run_query(
                    host, port, service, username, password,
                    date_from, date_to,
                    primary_id_val=primary_id_val.strip() or None,
                    id1_val=id1_val.strip() or None,
                    id2_val=id2_val.strip() or None,
                )
                st.session_state["excel_buf"] = build_excel(st.session_state["results"])
                filters = []
                if primary_id_val.strip(): filters.append(f"{primary_id_label}={primary_id_val.strip()}")
                if id1_val.strip(): filters.append(f"{id1_label}={id1_val.strip()}")
                if id2_val.strip(): filters.append(f"{id2_label}={id2_val.strip()}")
                st.session_state["query_range"] = (
                    f"{date_from}  →  {date_to}"
                    + (f"  |  {', '.join(filters)}" if filters else "")
                )
            except Exception as e:
                st.error(f"Query failed: {e}")
                st.session_state.pop("results", None)

# ── Results ────────────────────────────────────────────────────────────────
if "results" in st.session_state:
    results = st.session_state["results"]

    total     = len(results)
    diff_count = sum(1 for r in results if any(s["has_diff"] for s in r["sets"]))

    st.caption(f"Range: {st.session_state.get('query_range', '')}")

    m1, m2, m3, m4 = st.columns([1, 1, 1, 1])
    m1.metric("Total Records", total)
    m2.metric("Records with Differences", diff_count)
    m3.metric("Identical Records", total - diff_count)
    with m4:
        fname = f"json_diff_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        m4.download_button(
            label="Export to Excel",
            data=st.session_state["excel_buf"],
            file_name=fname,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

    # Filter
    q = st.text_input("Filter by ID, path, or value", placeholder="Search...")
    if q:
        ql = q.lower()
        results = [
            r for r in results
            if ql in r["id"].lower()
            or ql in r["dat_time"].lower()
            or any(ql in v.lower() for v in r["extra"].values())
            or any(
                ql in d["path"].lower() or ql in d["old"].lower() or ql in d["new"].lower()
                for s in r["sets"] for d in s["diffs"]
            )
        ]
        st.caption(f"{len(results)} matching records")

    st.divider()

    extra_ids = COL.get("extra_ids", [])

    for rec in results:
        has_any_diff = any(s["has_diff"] for s in rec["sets"])
        label_color  = "#8b0000" if has_any_diff else "#276221"
        diff_label   = f"{'⚠ Differences found' if has_any_diff else '✓ Identical'}"

        with st.expander(
            f"**{COL['primary_id']}: {rec['id']}** &nbsp;|&nbsp; {rec['dat_time']}"
            + (f"  —  " + "  |  ".join(f"{k}: {rec['extra'][k]}" for k in extra_ids if rec['extra'].get(k)) if extra_ids else ""),
            expanded=has_any_diff,
        ):
            for s in rec["sets"]:
                st.markdown(
                    f"**{s['label']}** &nbsp; `{s['left_label']}` vs `{s['right_label']}`",
                    unsafe_allow_html=True,
                )
                if not s["has_diff"]:
                    st.success("No differences")
                else:
                    render_diffs(s["diffs"])
                st.markdown("---")

